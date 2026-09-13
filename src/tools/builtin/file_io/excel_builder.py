"""Build formatted Excel workbooks from structured sheet data."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .office_errors import DocumentBuildError
from .office_presets import EXCEL_PRESETS, ExcelPreset
from .office_support import require_pkg


def _require_openpyxl() -> None:
    require_pkg(
        "openpyxl",
        DocumentBuildError,
        "未安装 openpyxl。请运行: pip install openpyxl",
        "openpyxl.styles",
    )


def _border_for_style(preset: ExcelPreset):
    from openpyxl.styles import Border, Side

    thin = Side(style="thin")
    if preset.border_style == "sanxian":
        medium = Side(style="medium")
        return Border(top=thin, bottom=medium, left=thin, right=thin)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _write_sheet(ws, sheet: dict[str, Any], preset: ExcelPreset) -> int:
    from openpyxl.styles import Alignment, Font

    row_idx = 1
    title = str(sheet.get("title", "") or "").strip()
    if title:
        ws.cell(row=row_idx, column=1, value=title)
        ws.cell(row=row_idx, column=1).font = Font(name=preset.title_font, size=preset.title_size_pt, bold=True)
        row_idx += 2

    headers = [str(h) for h in sheet.get("headers", [])]
    rows = sheet.get("rows", [])
    if not headers and rows:
        headers = [f"列{i + 1}" for i in range(len(rows[0]))]

    border = _border_for_style(preset)
    if headers:
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=header)
            cell.font = Font(name=preset.header_font, size=preset.header_size_pt, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        row_idx += 1

    body_rows = 0
    for row in rows:
        if not isinstance(row, (list, tuple)):
            continue
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name=preset.body_font, size=preset.body_size_pt)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
        row_idx += 1
        body_rows += 1

    for col_idx in range(1, max(len(headers), 1) + 1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = 18
    return body_rows


def build_excel_workbook(
    output_path: Path,
    *,
    sheets: list[dict[str, Any]],
    style: str = "generic",
) -> dict[str, Any]:
    """Create an .xlsx file from structured sheet definitions."""
    _require_openpyxl()
    from openpyxl import Workbook

    preset_key = style.strip().lower() or "generic"
    if preset_key not in EXCEL_PRESETS:
        raise DocumentBuildError(f"未知的 Excel 样式 '{style}'。可用: generic, official, sanxian。")
    preset = EXCEL_PRESETS[preset_key]

    if not sheets:
        raise DocumentBuildError("至少需要一个工作表。")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default_ws = workbook.active
    workbook.remove(default_ws)

    total_rows = 0
    for index, sheet in enumerate(sheets):
        name = str(sheet.get("name") or f"Sheet{index + 1}")[:31]
        ws = workbook.create_sheet(title=name)
        total_rows += _write_sheet(ws, sheet, preset)

    workbook.save(str(output_path))
    return {
        "path": str(output_path.resolve()),
        "style": preset_key,
        "sheets": len(sheets),
        "data_rows": total_rows,
    }
