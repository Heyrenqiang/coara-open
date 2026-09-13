"""Extract text from Excel (.xlsx) workbooks for the read tool."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .office_errors import DocumentReadError
from .office_support import require_pkg


def _require_openpyxl() -> None:
    require_pkg("openpyxl", DocumentReadError, "未安装 openpyxl。请运行: pip install openpyxl")


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def _render_sheet_as_markdown(ws) -> tuple[str, int]:
    """Render one worksheet as markdown. Returns (text, body_row_count).

    The first non-empty row is treated as the header. Returns ("", 0) for empty sheets.
    """
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        str_row = [_cell_to_str(cell) for cell in row]
        if any(cell for cell in str_row):
            rows.append(str_row)

    if not rows:
        return "", 0

    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]

    header = rows[0]
    lines = ["| " + " | ".join(header) + " |"]
    lines.append("| " + " | ".join("---" for _ in range(width)) + " |")
    for body_row in rows[1:]:
        lines.append("| " + " | ".join(body_row) + " |")
    return "\n".join(lines), max(0, len(rows) - 1)


def read_excel_workbook(path: Path) -> tuple[str, dict[str, Any]]:
    """Extract text from a .xlsx file → (text, meta).

    Each worksheet is rendered as a markdown table under a ``### Sheet: <name>`` heading.
    ``data_only=True`` returns cached computed values rather than formulas. Output is
    plain text suitable for the read tool's slice / line-number formatting.

    Raises:
        DocumentReadError: if openpyxl is missing or the file cannot be parsed.
    """
    _require_openpyxl()
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentReadError(f"无法解析 Excel 工作簿: {path}（{exc}）") from exc

    parts: list[str] = []
    total_rows = 0
    sheet_count = len(workbook.worksheets)
    try:
        for ws in workbook.worksheets:
            rendered, body_rows = _render_sheet_as_markdown(ws)
            total_rows += body_rows
            if not rendered:
                continue
            parts.append(f"### Sheet: {ws.title}")
            parts.append("")
            parts.append(rendered)
            parts.append("")
    finally:
        workbook.close()

    content = "\n".join(parts).strip()
    meta: dict[str, Any] = {
        "format": "xlsx",
        "engine": "openpyxl",
        "sheets": sheet_count,
        "rows": total_rows,
    }
    return content, meta
