"""Extract text from PowerPoint (.pptx) presentations for the read tool."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from src.core.text import render_table_as_markdown

from .office_errors import DocumentReadError


def _require_pptx() -> None:
    try:
        from pptx import Presentation  # noqa: F401
    except ImportError as exc:  # pragma: no cover - exercised via missing-dep test
        raise DocumentReadError("未安装 python-pptx。请运行: pip install python-pptx") from exc


def read_pptx_presentation(path: Path) -> tuple[str, dict[str, Any]]:
    """Extract text from a .pptx file → (text, meta).

    Each slide is rendered under a ``## 幻灯片 N`` heading, followed by the text of its
    shapes (text frames and tables) and any speaker notes. Output is plain text suitable
    for the read tool's slice / line-number formatting.

    Raises:
        DocumentReadError: if python-pptx is missing or the file cannot be parsed.
    """
    _require_pptx()
    from pptx import Presentation

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise DocumentReadError(f"无法解析 PowerPoint 演示文稿: {path}（{exc}）") from exc

    parts: list[str] = []
    slide_count = len(presentation.slides)

    for index, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = [f"## 幻灯片 {index}"]
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                slide_parts.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False):
                rendered = render_table_as_markdown(shape.table)
                if rendered:
                    slide_parts.append("")
                    slide_parts.append(rendered)
                    slide_parts.append("")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_parts.append(f"备注: {notes}")
        parts.append("\n".join(slide_parts))

    content = "\n\n".join(parts).strip()
    meta: dict[str, Any] = {
        "format": "pptx",
        "engine": "python-pptx",
        "slides": slide_count,
    }
    return content, meta
