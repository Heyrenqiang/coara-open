"""Build formatted Word documents from Markdown-like text."""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .office_errors import DocumentBuildError
from .office_presets import WORD_PRESETS, WordPreset
from .office_support import atomic_output_path, require_pkg

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_OFFICIAL_H1_RE = re.compile(r"^[一二三四五六七八九十百零〇]+、\s*(.+)$")
_OFFICIAL_H2_RE = re.compile(r"^（[一二三四五六七八九十百零〇]+）\s*(.+)$")
_OFFICIAL_H3_RE = re.compile(r"^\d+\.\s+(.+)$")
_OFFICIAL_H4_RE = re.compile(r"^（\d+）\s*(.+)$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_ORDERED_RE = re.compile(r"^\d+\.\s+(.*)$")
_INLINE_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_INLINE_UNDERLINE_BOLD_RE = re.compile(r"__(.+?)__")
_INLINE_SEGMENT_RE = re.compile(r"\*\*(.+?)\*\*|__(.+?)__")


@dataclass(frozen=True, slots=True)
class _WordBuildContext:
    preset: WordPreset
    from_template: bool

    @property
    def apply_preset_typography(self) -> bool:
        return not self.from_template


def _strip_inline_md(text: str) -> str:
    """Remove inline markers for line classification only."""
    text = _INLINE_BOLD_RE.sub(r"\1", text)
    return _INLINE_UNDERLINE_BOLD_RE.sub(r"\1", text)


def _iter_inline_segments(text: str) -> list[tuple[str, bool]]:
    """Split text into (segment, is_bold) tuples."""
    segments: list[tuple[str, bool]] = []
    pos = 0
    for match in _INLINE_SEGMENT_RE.finditer(text):
        if match.start() > pos:
            segments.append((text[pos : match.start()], False))
        bold_text = match.group(1) or match.group(2) or ""
        segments.append((bold_text, True))
        pos = match.end()
    if pos < len(text):
        segments.append((text[pos:], False))
    if not segments:
        segments.append((text, False))
    return segments


def _add_inline_runs(
    paragraph,
    text: str,
    ctx: _WordBuildContext,
    *,
    font: str,
    size: float,
    default_bold: bool = False,
) -> None:
    stripped = text.strip()
    if not stripped:
        return
    apply_typography = ctx.apply_preset_typography
    for segment, is_bold in _iter_inline_segments(stripped):
        if not segment:
            continue
        run = paragraph.add_run(segment)
        _set_east_asia_font(
            run,
            font,
            size,
            bold=default_bold or is_bold,
            apply_typography=apply_typography,
        )


def _require_docx() -> None:
    require_pkg(
        "docx",
        DocumentBuildError,
        "未安装 python-docx。请运行: pip install python-docx",
        "docx.enum.text",
        "docx.oxml.ns",
        "docx.shared",
    )


def _set_east_asia_font(
    run,
    font_name: str,
    size_pt: float,
    *,
    bold: bool = False,
    apply_typography: bool = True,
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    if apply_typography:
        run.font.name = font_name
        run.font.size = Pt(size_pt)
    if bold:
        run.font.bold = bold
    if not apply_typography:
        return

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        from docx.oxml import OxmlElement

        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)


def _configure_page(section, preset: WordPreset) -> None:
    from docx.shared import Mm

    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(preset.top_margin_mm)
    section.bottom_margin = Mm(preset.bottom_margin_mm)
    section.left_margin = Mm(preset.left_margin_mm)
    section.right_margin = Mm(preset.right_margin_mm)


def _apply_paragraph_style(
    paragraph,
    ctx: _WordBuildContext,
    *,
    font: str,
    size: float,
    force_bold: bool | None = None,
) -> None:
    if not ctx.apply_preset_typography:
        return

    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    preset = ctx.preset
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(preset.line_spacing_pt)
    fmt.space_after = Pt(0)
    if not paragraph.text.strip():
        return

    for run in paragraph.runs:
        run_bold = force_bold if force_bold is not None else bool(run.font.bold)
        _set_east_asia_font(run, font, size, bold=run_bold, apply_typography=True)


def _heading_font(preset: WordPreset, level: int) -> tuple[str, float]:
    if level <= 1:
        return preset.h1_font, preset.h1_size_pt
    if level == 2:
        return preset.h2_font, preset.h2_size_pt
    if level == 3:
        return preset.h3_font, preset.h3_size_pt
    return preset.h4_font, preset.h4_size_pt


def _add_title(document, title: str, ctx: _WordBuildContext) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    preset = ctx.preset
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title.strip())
    _set_east_asia_font(
        run,
        preset.title_font,
        preset.title_size_pt,
        bold=True,
        apply_typography=ctx.apply_preset_typography,
    )
    _apply_paragraph_style(
        paragraph,
        ctx,
        font=preset.title_font,
        size=preset.title_size_pt,
        force_bold=True,
    )


def _add_body_paragraph(document, text: str, ctx: _WordBuildContext) -> None:
    from docx.shared import Pt

    preset = ctx.preset
    paragraph = document.add_paragraph()
    _add_inline_runs(paragraph, text, ctx, font=preset.body_font, size=preset.body_size_pt)
    if ctx.apply_preset_typography and preset.first_line_indent_chars > 0 and text.strip():
        paragraph.paragraph_format.first_line_indent = Pt(preset.body_size_pt * preset.first_line_indent_chars)
    _apply_paragraph_style(paragraph, ctx, font=preset.body_font, size=preset.body_size_pt)


def _add_heading(document, level: int, text: str, ctx: _WordBuildContext) -> None:
    font, size = _heading_font(ctx.preset, level)
    paragraph = document.add_paragraph()
    _add_inline_runs(paragraph, text, ctx, font=font, size=size, default_bold=True)
    _apply_paragraph_style(paragraph, ctx, font=font, size=size, force_bold=True)


def _add_bullet_paragraph(document, text: str) -> Any:
    try:
        return document.add_paragraph(style="List Bullet")
    except KeyError:
        paragraph = document.add_paragraph()
        paragraph.add_run("• ")
        return paragraph


def _add_bullet(document, text: str, ctx: _WordBuildContext) -> None:
    preset = ctx.preset
    paragraph = _add_bullet_paragraph(document, text)
    _add_inline_runs(paragraph, text, ctx, font=preset.body_font, size=preset.body_size_pt)
    _apply_paragraph_style(paragraph, ctx, font=preset.body_font, size=preset.body_size_pt)


def _classify_line(line: str, *, official_mode: bool) -> tuple[str, str] | None:
    """Return (kind, text) for structured lines, or None for body."""
    stripped = line.strip()
    if not stripped:
        return None
    bare = _strip_inline_md(stripped)

    if official_mode:
        for pattern, kind in (
            (_OFFICIAL_H1_RE, "h1"),
            (_OFFICIAL_H2_RE, "h2"),
            (_OFFICIAL_H3_RE, "h3"),
            (_OFFICIAL_H4_RE, "h4"),
        ):
            match = pattern.match(bare)
            if match:
                return kind, match.group(1)

    heading = _HEADING_RE.match(stripped)
    if heading:
        level = min(len(heading.group(1)), 6)
        return f"h{level}", heading.group(2)

    bullet = _BULLET_RE.match(stripped)
    if bullet:
        return "bullet", bullet.group(1)

    if not official_mode:
        ordered = _ORDERED_RE.match(stripped)
        if ordered:
            return "body", ordered.group(1)

    return None


def _render_markdown_body(document, content: str, ctx: _WordBuildContext, *, style: str) -> None:
    official_mode = style.strip().lower() == "official"
    lines = content.replace("\r\n", "\n").split("\n")
    for raw_line in lines:
        line = raw_line.rstrip()
        if not line.strip():
            continue
        classified = _classify_line(line, official_mode=official_mode)
        if classified is None:
            _add_body_paragraph(document, line, ctx)
            continue
        kind, text = classified
        if kind == "bullet":
            _add_bullet(document, text, ctx)
            continue
        if kind.startswith("h"):
            level = int(kind[1:])
            _add_heading(document, level, text, ctx)
            continue
        _add_body_paragraph(document, text, ctx)


def build_word_document(
    output_path: Path,
    *,
    content: str,
    style: str = "report",
    title: str | None = None,
    template_path: Path | None = None,
) -> dict[str, Any]:
    """Create a .docx file from Markdown-like text using a built-in preset."""
    _require_docx()
    from docx import Document

    preset_key = style.strip().lower() or "report"
    if preset_key not in WORD_PRESETS:
        raise DocumentBuildError(f"未知的 Word 样式 '{style}'。可用: report, official, academic。")
    preset = WORD_PRESETS[preset_key]
    from_template = template_path is not None
    ctx = _WordBuildContext(preset=preset, from_template=from_template)

    with atomic_output_path(output_path) as temp_path:
        if template_path is not None:
            if not template_path.is_file():
                raise DocumentBuildError(f"未找到 Word 模板: {template_path}")
            shutil.copy2(template_path, temp_path)
            document = Document(str(temp_path))
        else:
            document = Document()
            for section in document.sections:
                _configure_page(section, preset)

        if title and title.strip():
            _add_title(document, title, ctx)

        body = content.strip()
        if not body:
            raise DocumentBuildError("文档内容不能为空。")
        _render_markdown_body(document, body, ctx, style=preset_key)

        document.save(str(temp_path))
    return {
        "path": str(output_path.resolve()),
        "style": preset_key,
        "paragraphs": len(document.paragraphs),
        "title": title or "",
        "from_template": from_template,
    }
