"""Extract text from Word (.docx) documents for the read tool."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from src.core.text import render_table_as_markdown

from .office_errors import DocumentReadError
from .office_support import require_pkg


def _require_docx() -> None:
    require_pkg("docx", DocumentReadError, "未安装 python-docx。请运行: pip install python-docx")


def _heading_prefix(style_name: str) -> str:
    """Map a Word style name to a markdown heading prefix, or '' for non-headings."""
    if not style_name:
        return ""
    name = style_name.strip()
    if name.lower().startswith("heading"):
        parts = name.split()
        if len(parts) == 2 and parts[1].isdigit():
            level = max(1, min(6, int(parts[1])))
            return "#" * level + " "
    if name.lower() == "title":
        return "# "
    return ""


def read_word_document(path: Path) -> tuple[str, dict[str, Any]]:
    """Extract text from a .docx file → (text, meta).

    Paragraphs are joined with newlines; Heading styles become markdown ``#`` prefixes.
    Tables are rendered as markdown tables. Output is plain text suitable for the
    read tool's existing slice / line-number formatting.

    Raises:
        DocumentReadError: if python-docx is missing or the file cannot be parsed.
    """
    _require_docx()
    from docx import Document

    try:
        document = Document(str(path))
    except Exception as exc:
        raise DocumentReadError(f"无法解析 Word 文档: {path}（{exc}）") from exc

    parts: list[str] = []
    paragraph_count = 0
    table_count = len(document.tables)

    # python-docx exposes paragraphs and tables as separate iterables; reading them in
    # document order requires walking the body XML. For LLM consumption a faithful
    # in-order weave is not essential — paragraphs first, then tables — and keeps the
    # implementation robust against mixed content.
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        prefix = _heading_prefix(para.style.name if para.style else "")
        parts.append(f"{prefix}{text}" if prefix else text)
        paragraph_count += 1

    for table in document.tables:
        rendered = render_table_as_markdown(table)
        if rendered:
            parts.append("")
            parts.append(rendered)
            parts.append("")

    content = "\n".join(parts).strip()
    meta: dict[str, Any] = {
        "format": "docx",
        "engine": "python-docx",
        "paragraphs": paragraph_count,
        "tables": table_count,
    }
    return content, meta
