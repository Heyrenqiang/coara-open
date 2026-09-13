"""Tests for file I/O tools (read, write, edit, glob, grep)."""

from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

from src.tools.builtin.file_io import grep as grep_module
from src.tools.builtin.file_io.edit import EditTool
from src.tools.builtin.file_io.glob import GlobTool
from src.tools.builtin.file_io.grep import GrepTool
from src.tools.builtin.file_io.read import ReadTool
from src.tools.builtin.file_io.search_support import GrepParams
from src.tools.builtin.file_io.write import WriteTool
from src.tools.cache import tool_cache


@pytest.fixture
def clear_read_cache() -> None:
    tool_cache.invalidate("read")
    yield
    tool_cache.invalidate("read")


@pytest.mark.asyncio
async def test_read_tool_decodes_cp1252_text(tmp_path: Path, clear_read_cache: None) -> None:
    target = tmp_path / "legacy.txt"
    target.write_bytes("Copyright \xa9 2024 All rights reserved".encode("cp1252"))
    tool = ReadTool(workspace_root=tmp_path)

    result = await tool.create_invocation({"path": str(target)}).execute()

    assert not result.is_error
    assert "Copyright © 2024 All rights reserved" in result.content


@pytest.mark.asyncio
async def test_read_tool_decodes_gbk_text(tmp_path: Path, clear_read_cache: None) -> None:
    target = tmp_path / "legacy_cn.txt"
    target.write_bytes("你好世界\r\n第二行".encode("gbk"))
    tool = ReadTool(workspace_root=tmp_path)

    result = await tool.create_invocation({"path": str(target)}).execute()

    assert not result.is_error
    assert "你好世界" in result.content
    assert "第二行" in result.content


@pytest.mark.asyncio
async def test_read_tool_rejects_true_binary(tmp_path: Path, clear_read_cache: None) -> None:
    target = tmp_path / "fake.txt"
    target.write_bytes(b"MZ\x90\x00\x03\x00\x00\x00" + bytes(range(256)) * 4)
    tool = ReadTool(workspace_root=tmp_path)

    result = await tool.create_invocation({"path": str(target)}).execute()

    assert result.is_error
    assert "二进制" in result.content


@pytest.mark.asyncio
async def test_edit_tool_roundtrips_legacy_encoding(tmp_path: Path, clear_read_cache: None) -> None:
    target = tmp_path / "legacy_cn.txt"
    original_bytes = "你好世界\r\n第二行".encode("gbk")
    target.write_bytes(original_bytes)
    shared_states: dict[str, object] = {}
    read_tool = ReadTool(read_state_store=shared_states, workspace_root=tmp_path)
    edit_tool = EditTool(read_state_store=shared_states, workspace_root=tmp_path)

    read_result = await read_tool.create_invocation({"path": str(target)}).execute()
    assert not read_result.is_error
    edit_result = await edit_tool.create_invocation(
        {"path": str(target), "old_string": "第二行", "new_string": "第三行"}
    ).execute()
    assert not edit_result.is_error

    assert target.read_bytes() == "你好世界\r\n第三行".encode("gbk")


@pytest.mark.asyncio
async def test_write_tool_overwrites_existing_file(tmp_path: Path) -> None:
    target = tmp_path / "note.txt"
    tool = WriteTool(workspace_root=tmp_path)

    inv1 = tool.create_invocation({"path": str(target), "contents": "alpha"})
    first = await inv1.execute()

    inv2 = tool.create_invocation({"path": str(target), "contents": "beta"})
    second = await inv2.execute()

    assert not first.is_error
    assert not second.is_error
    assert target.read_text(encoding="utf-8") == "beta"


@pytest.mark.asyncio
async def test_read_tool_reuses_cached_result_for_unchanged_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    clear_read_cache: None,
) -> None:
    target = tmp_path / "note.txt"
    target.write_text("alpha", encoding="utf-8")
    shared_states: dict[str, object] = {}
    tool = ReadTool(read_state_store=shared_states, workspace_root=tmp_path)

    first = await tool.create_invocation({"path": str(target)}).execute()
    assert not first.is_error

    def fail_read_text_file(_: Path) -> tuple[str, object]:
        raise AssertionError("read_text_file should not be called for unchanged cached reads")

    monkeypatch.setattr("src.tools.builtin.file_io.read.read_text_file", fail_read_text_file)

    second = await tool.create_invocation({"path": str(target)}).execute()
    assert not second.is_error
    assert second.content == first.content


@pytest.mark.asyncio
async def test_write_tool_reuses_snapshot_after_successful_write(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    clear_read_cache: None,
) -> None:
    target = tmp_path / "note.txt"
    target.write_text("alpha", encoding="utf-8")
    shared_states: dict[str, object] = {}
    read_tool = ReadTool(read_state_store=shared_states, workspace_root=tmp_path)
    write_tool = WriteTool(read_state_store=shared_states, workspace_root=tmp_path)

    first_read = await read_tool.create_invocation({"path": str(target)}).execute()
    assert not first_read.is_error

    first_write = await write_tool.create_invocation({"path": str(target), "contents": "beta"}).execute()
    assert not first_write.is_error

    def fail_read_text_file(_: Path) -> tuple[str, object]:
        raise AssertionError("write should reuse the fresh snapshot instead of reading from disk")

    monkeypatch.setattr("src.tools.builtin.file_io.file_support.read_text_file", fail_read_text_file)

    second_write = await write_tool.create_invocation({"path": str(target), "contents": "gamma"}).execute()

    assert not second_write.is_error
    assert target.read_text(encoding="utf-8") == "gamma"


@pytest.mark.asyncio
async def test_file_tools_use_explicit_workspace_root_instead_of_cwd(tmp_path: Path, clear_read_cache: None) -> None:
    previous_cwd = Path.cwd()
    workspace = tmp_path / "workspace"
    elsewhere = tmp_path / "elsewhere"
    workspace.mkdir()
    elsewhere.mkdir()
    os.chdir(elsewhere)
    try:
        target = workspace / "note.txt"
        outside = elsewhere / "outside.txt"
        target.write_text("alpha", encoding="utf-8")
        outside.write_text("beta", encoding="utf-8")

        shared_states: dict[str, object] = {}
        read_tool = ReadTool(read_state_store=shared_states, workspace_root=workspace)
        write_tool = WriteTool(read_state_store=shared_states, workspace_root=workspace)

        inside_result = await read_tool.create_invocation({"path": str(target)}).execute()
        outside_result = await read_tool.create_invocation({"path": str(outside)}).execute()
        outside_write = await write_tool.create_invocation({"path": str(outside), "contents": "gamma"}).execute()
        write_result = await write_tool.create_invocation({"path": str(target), "contents": "gamma"}).execute()

        assert not inside_result.is_error
        assert not outside_result.is_error
        # 越界写入不再硬拒：直接 execute 放行（审批门在执行层 requires_approval）。
        assert not outside_write.is_error
        assert outside.read_text(encoding="utf-8") == "gamma"
        assert write_tool.requires_approval({"path": str(outside), "contents": "x"}) is True
        assert write_tool.requires_approval({"path": str(target), "contents": "x"}) is False
        assert not write_result.is_error
        assert target.read_text(encoding="utf-8") == "gamma"
    finally:
        os.chdir(previous_cwd)


@pytest.mark.asyncio
async def test_glob_tool_sorts_alphabetically(tmp_path: Path) -> None:
    (tmp_path / "alpha.txt").write_text("a", encoding="utf-8")
    (tmp_path / "beta.txt").write_text("b", encoding="utf-8")
    tool = GlobTool(workspace_root=tmp_path)

    files = await tool.create_invocation({"path": str(tmp_path), "pattern": "*.txt"}).execute()

    assert not files.is_error
    assert files.content == f"{(tmp_path / 'alpha.txt').as_posix()}\n{(tmp_path / 'beta.txt').as_posix()}"


@pytest.mark.asyncio
async def test_glob_tool_rejects_missing_path(tmp_path: Path) -> None:
    missing = tmp_path / "missing"
    result = (
        await GlobTool(workspace_root=tmp_path).create_invocation({"path": str(missing), "pattern": "*.txt"}).execute()
    )

    assert result.is_error
    assert "路径不存在" in (result.content or "")


@pytest.mark.asyncio
async def test_glob_tool_rejects_relative_path(tmp_path: Path) -> None:
    result = await GlobTool(workspace_root=tmp_path).create_invocation({"path": "src", "pattern": "*.py"}).execute()

    assert result.is_error
    assert "绝对路径" in (result.content or "")


@pytest.mark.asyncio
async def test_glob_tool_skips_node_modules(tmp_path: Path) -> None:
    (tmp_path / "src").mkdir()
    (tmp_path / "src" / "app.py").write_text("x", encoding="utf-8")
    (tmp_path / "node_modules").mkdir()
    (tmp_path / "node_modules" / "pkg.js").write_text("y", encoding="utf-8")
    tool = GlobTool(workspace_root=tmp_path)

    files = await tool.create_invocation({"path": str(tmp_path), "pattern": "*.py"}).execute()
    assert not files.is_error
    assert files.content == (tmp_path / "src" / "app.py").as_posix()
    assert files.metadata["count"] == 1


@pytest.mark.asyncio
async def test_grep_output_modes_and_offset_python_fallback(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(grep_module, "find_rg_executable", lambda: None)

    (tmp_path / "a.py").write_text("def foo():\n    return 1\n", encoding="utf-8")
    (tmp_path / "b.py").write_text("class Bar:\n    pass\n", encoding="utf-8")
    from src.tools.builtin.file_io.search_support import MAX_GREP_FILE_BYTES

    (tmp_path / "big.py").write_text("x" * (MAX_GREP_FILE_BYTES + 1), encoding="utf-8")
    tool = GrepTool(workspace_root=tmp_path)

    files = await tool.create_invocation(
        {"pattern": "def ", "output_mode": "files_with_matches", "type": "py"}
    ).execute()
    assert not files.is_error
    assert files.content == (tmp_path / "a.py").as_posix()

    counts = await tool.create_invocation({"pattern": "def ", "output_mode": "count", "glob": "*.py"}).execute()
    assert not counts.is_error
    assert counts.content == f"{(tmp_path / 'a.py').as_posix()}:1"

    content = await tool.create_invocation(
        {
            "pattern": "def ",
            "output_mode": "content",
            "context_lines": 1,
            "limit": 3,
        }
    ).execute()
    assert not content.is_error
    assert f"{(tmp_path / 'a.py').as_posix()}:1:" in (content.content or "")
    assert f"{(tmp_path / 'a.py').as_posix()}-2-" in (content.content or "")

    paged = await tool.create_invocation(
        {"pattern": ".", "glob": "a.py", "output_mode": "content", "limit": 1, "offset": 0}
    ).execute()
    assert not paged.is_error
    assert paged.metadata["truncated"] is True

    skipped = await tool.create_invocation({"pattern": "x", "glob": "big.py"}).execute()
    assert not skipped.is_error
    assert skipped.metadata.get("skipped_large_files") == [(tmp_path / "big.py").as_posix()]


@pytest.mark.asyncio
async def test_grep_rg_dialect_error_auto_fixes_and_hints(tmp_path: Path) -> None:
    """ripgrep 正则解析失败后自动转义重试；非解析类错误仍报错并带恢复提示。"""
    if grep_module.find_rg_executable() is None:
        pytest.skip("ripgrep not available")

    (tmp_path / "a.py").write_text('f"{ = 1\n', encoding="utf-8")
    tool = GrepTool(workspace_root=tmp_path)

    # 字面 { 在 Rust 正则里是量词起始符 → regex parse error → 自动转义重试成功
    bad_brace = await tool.create_invocation({"pattern": 'f"{', "output_mode": "content"}).execute()
    assert not bad_brace.is_error
    assert "自动转义" in (bad_brace.metadata.get("note") or "")

    # 字面换行在 rg 默认逐行模式下非法 → 非解析类错误，仍报错带提示
    bad_newline = await tool.create_invocation({"pattern": "x = 1\ny = 2"}).execute()
    assert bad_newline.is_error
    assert "Rust 正则" in (bad_newline.content or "")


@pytest.mark.asyncio
async def test_grep_accepts_single_file_path(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(grep_module, "find_rg_executable", lambda: None)

    report = tmp_path / "report.md"
    report.write_text("执行摘要\n信丰西站\n", encoding="utf-8")
    tool = GrepTool(workspace_root=tmp_path)

    result = await tool.create_invocation(
        {"pattern": "执行摘要", "path": str(report), "output_mode": "content"}
    ).execute()
    assert not result.is_error
    assert "执行摘要" in (result.content or "")


def test_grep_accepts_cursor_aliases() -> None:
    params = GrepParams.from_tool_params({"pattern": "foo", "glob_file_search": "*.py", "-i": False})
    assert params.glob == "*.py"
    assert params.case_sensitive is True


@pytest.mark.asyncio
async def test_edit_tool_attaches_diff_display(tmp_path: Path) -> None:
    target = tmp_path / "code.py"
    target.write_text("foo = 1\nbar = 2\n", encoding="utf-8")
    tool = EditTool(read_state_store={}, workspace_root=tmp_path)
    result = await tool.create_invocation(
        {"path": str(target), "old_string": "foo = 1", "new_string": "foo = 42"}
    ).execute()
    assert not result.is_error
    assert result.display
    assert result.metadata.get("lines_added", 0) >= 1


@pytest.mark.asyncio
async def test_read_tool_returns_line_number_prefix(tmp_path: Path) -> None:
    target = tmp_path / "note.txt"
    target.write_text("line1\nline2\nline3\n", encoding="utf-8")
    tool = ReadTool(workspace_root=tmp_path)
    result = await tool.create_invocation({"path": str(target), "offset": 2, "limit": 1}).execute()
    assert not result.is_error
    assert result.content == "2|line2"


@pytest.mark.asyncio
async def test_write_tool_has_no_diff_display(tmp_path: Path) -> None:
    target = tmp_path / "note.txt"
    target.write_text("hello\n", encoding="utf-8")
    tool = WriteTool(read_state_store={}, workspace_root=tmp_path)
    result = await tool.create_invocation({"path": str(target), "contents": "hello world\n"}).execute()
    assert not result.is_error
    assert not result.display
    assert "lines_added" not in result.metadata


def test_edit_logic_matches_execute(tmp_path: Path) -> None:
    from src.tools.builtin.file_io.edit_logic import simulate_edit_replacement

    content = "foo = 1\nbar = 2\n"
    simulated = simulate_edit_replacement(content, "foo = 1", "foo = 42")
    assert simulated.ok
    assert simulated.matches == 1
    assert "foo = 42" in simulated.new_content


def test_edit_logic_fuzzy_trailing_whitespace_preserves_file_indent() -> None:
    """rstrip 模糊：模型多抄了行尾空格时，仍用文件原文锚定再替换。"""
    from src.tools.builtin.file_io.edit_logic import simulate_edit_replacement

    content = "    foo = 1\nbar = 2\n"
    # 模型多带了行尾空格 → 精确子串失败，走 rstrip 级；替换目标仍是文件中的「    foo = 1」
    simulated = simulate_edit_replacement(content, "    foo = 1  ", "    foo = 42")
    assert simulated.ok
    assert simulated.new_content == "    foo = 42\nbar = 2\n"


def test_edit_logic_fuzzy_unicode_dash_and_quotes() -> None:
    from src.tools.builtin.file_io.edit_logic import simulate_edit_replacement

    content = "note = “hello—world”\n"
    simulated = simulate_edit_replacement(content, 'note = "hello-world"', 'note = "ok"')
    assert simulated.ok
    assert simulated.new_content == 'note = "ok"\n'


def test_edit_logic_fuzzy_ambiguous_different_indents_fails() -> None:
    """两处标点不同、Unicode 归一后相同 → 不擅自挑选。"""
    from src.tools.builtin.file_io.edit_logic import simulate_edit_replacement

    content = "a—b\na–b\n"  # em-dash vs en-dash
    simulated = simulate_edit_replacement(content, "a-b", "a=b")
    assert not simulated.ok
    assert simulated.error == "未找到匹配文本"


def test_edit_logic_exact_still_requires_unique() -> None:
    from src.tools.builtin.file_io.edit_logic import simulate_edit_replacement

    content = "x = 1\nx = 1\n"
    simulated = simulate_edit_replacement(content, "x = 1", "x = 2")
    assert not simulated.ok
    assert "2 处匹配" in simulated.error


@pytest.mark.asyncio
async def test_read_tool_auto_limits_large_file(tmp_path: Path, clear_read_cache: None) -> None:
    from src.core.read_format import DEFAULT_FILE_READ_LIMIT_LINES

    target = tmp_path / "big.txt"
    total = DEFAULT_FILE_READ_LIMIT_LINES + 50
    target.write_text("\n".join(f"line-{i}" for i in range(1, total + 1)) + "\n", encoding="utf-8")
    tool = ReadTool(workspace_root=tmp_path)

    result = await tool.create_invocation({"path": str(target)}).execute()
    assert not result.is_error
    assert result.metadata.get("auto_limited") is True
    assert result.metadata.get("total_lines") == total
    assert f"仅显示前 {DEFAULT_FILE_READ_LIMIT_LINES}/{total} 行" in (result.content or "")
    assert "   1|line-1" in (result.content or "")
    assert f"{DEFAULT_FILE_READ_LIMIT_LINES}|line-{DEFAULT_FILE_READ_LIMIT_LINES}" in (result.content or "")
    assert f"line-{total}" not in (result.content or "")

    # 显式 limit 关闭默认上限语义（读完全部）
    full = await tool.create_invocation({"path": str(target), "limit": total}).execute()
    assert not full.is_error
    assert full.metadata.get("auto_limited") is not True
    assert f"line-{total}" in (full.content or "")


@pytest.mark.asyncio
async def test_glob_walk_budget_truncates(tmp_path: Path) -> None:
    from src.tools.builtin.file_io import search_support as ss

    small = ss.WalkBudget(max_depth=64, max_directories=10_000, max_entries=8)

    nested = tmp_path
    for i in range(20):
        nested = nested / f"d{i}"
        nested.mkdir()
        (nested / f"f{i}.txt").write_text("x", encoding="utf-8")

    status = ss.WalkStatus()
    paths = list(ss.iter_glob_paths(tmp_path, "*.txt", kind="file", budget=small, status=status))
    assert status.truncated is True
    assert len(paths) < 20


@pytest.mark.asyncio
async def test_grep_python_walk_budget_sets_metadata(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(grep_module, "find_rg_executable", lambda: None)
    from src.tools.builtin.file_io import search_support as ss

    real_budgeted = ss.iter_budgeted_walk

    def limited_walk(root, *, budget=None, status=None):
        return real_budgeted(
            root,
            budget=ss.WalkBudget(max_depth=64, max_directories=10_000, max_entries=6),
            status=status,
        )

    monkeypatch.setattr(grep_module, "iter_budgeted_walk", limited_walk)

    nested = tmp_path
    for i in range(15):
        nested = nested / f"d{i}"
        nested.mkdir()
        (nested / "hit.py").write_text("needle\n", encoding="utf-8")

    tool = GrepTool(workspace_root=tmp_path)
    result = await tool.create_invocation({"pattern": "needle", "output_mode": "files_with_matches"}).execute()
    assert not result.is_error
    assert result.metadata.get("walk_truncated") is True
    assert "遍历已达预算上限" in (result.content or "")


@pytest.mark.asyncio
async def test_read_tool_reads_docx_as_markdown(tmp_path: Path) -> None:
    from docx import Document

    target = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading("第一章", level=1)
    doc.add_paragraph("人工智能正在改变世界。")
    doc.add_paragraph("这是第二段。")
    doc.save(str(target))

    tool = ReadTool(workspace_root=tmp_path)
    result = await tool.create_invocation({"path": str(target)}).execute()

    assert not result.is_error
    assert "# 第一章" in (result.content or "")
    assert "人工智能正在改变世界" in (result.content or "")
    assert result.metadata.get("format") == "docx"


@pytest.mark.asyncio
async def test_edit_tool_edits_docx_paragraph(tmp_path: Path) -> None:
    from docx import Document

    target = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("人工智能正在改变世界。")
    doc.add_paragraph("这是第二段。")
    doc.save(str(target))

    tool = EditTool(read_state_store={}, workspace_root=tmp_path)
    result = await tool.create_invocation(
        {"path": str(target), "old_string": "人工智能正在改变世界", "new_string": "机器学习正在改变世界"}
    ).execute()

    assert not result.is_error
    assert result.metadata.get("matches_found") == 1
    assert result.display, "Office edit must attach terminal diff like text edit"
    joined_old = "\n".join(b.old_text for b in result.display)
    joined_new = "\n".join(b.new_text for b in result.display)
    assert "人工智能正在改变世界" in joined_old
    assert "机器学习正在改变世界" in joined_new

    edited_doc = Document(str(target))
    texts = [p.text for p in edited_doc.paragraphs]
    assert "机器学习正在改变世界。" in texts


@pytest.mark.asyncio
async def test_edit_tool_edits_xlsx_cell_with_diff(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")

    target = tmp_path / "sheet.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "hello world"
    ws["B1"] = "keep"
    wb.save(str(target))

    tool = EditTool(read_state_store={}, workspace_root=tmp_path)
    result = await tool.create_invocation(
        {"path": str(target), "old_string": "hello", "new_string": "hola"}
    ).execute()

    assert not result.is_error
    assert result.metadata.get("matches_found") == 1
    assert result.display
    assert any("hello" in (b.old_text or "") for b in result.display)
    assert any("hola" in (b.new_text or "") for b in result.display)

    wb2 = openpyxl.load_workbook(str(target))
    assert wb2.active["A1"].value == "hola world"


@pytest.mark.asyncio
async def test_edit_tool_edits_pptx_shape_with_diff(tmp_path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    target = tmp_path / "deck.pptx"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.paragraphs[0].text = "Draft title here"
    prs.save(str(target))

    tool = EditTool(read_state_store={}, workspace_root=tmp_path)
    result = await tool.create_invocation(
        {"path": str(target), "old_string": "Draft title", "new_string": "Final title"}
    ).execute()

    assert not result.is_error
    assert result.metadata.get("matches_found") == 1
    assert result.display
    assert any("Draft title" in (b.old_text or "") for b in result.display)
    assert any("Final title" in (b.new_text or "") for b in result.display)

    prs2 = pptx.Presentation(str(target))
    texts = [shape.text_frame.text for shape in prs2.slides[0].shapes if shape.has_text_frame]
    assert any("Final title here" in t for t in texts)


@pytest.mark.asyncio
async def test_edit_tool_returns_error_when_docx_text_not_found(tmp_path: Path) -> None:
    from docx import Document

    target = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("人工智能正在改变世界。")
    doc.save(str(target))

    tool = EditTool(read_state_store={}, workspace_root=tmp_path)
    result = await tool.create_invocation(
        {"path": str(target), "old_string": "不存在的文本", "new_string": "新文本"}
    ).execute()

    assert result.is_error
    assert "未找到要替换的文本" in (result.content or "")


@pytest.mark.asyncio
async def test_read_tool_extracts_pdf_text(tmp_path: Path) -> None:
    fitz = pytest.importorskip("fitz")

    target = tmp_path / "paper.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello Spatiotemporal Composability")
    doc.save(str(target))
    doc.close()

    tool = ReadTool(read_state_store={}, workspace_root=tmp_path)
    result = await tool.create_invocation({"path": str(target)}).execute()

    assert not result.is_error
    assert "PAGE 1" in (result.content or "")
    assert "Hello Spatiotemporal Composability" in (result.content or "")
    assert result.metadata.get("format") == "pdf"
    assert result.metadata.get("pages") == 1


@pytest.mark.asyncio
async def test_read_tool_pdf_without_pymupdf_gives_install_hint(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setitem(sys.modules, "fitz", None)

    target = tmp_path / "paper.pdf"
    target.write_bytes(b"%PDF-1.4 fake")

    tool = ReadTool(read_state_store={}, workspace_root=tmp_path)
    result = await tool.create_invocation({"path": str(target)}).execute()

    assert result.is_error
    assert "pip install pymupdf" in (result.content or "")
