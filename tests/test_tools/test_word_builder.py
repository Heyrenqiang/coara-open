from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from docx import Document

from src.tools.builtin.file_io.office_presets import WORD_PRESETS
from src.tools.builtin.file_io.word_builder import _heading_font, build_word_document


@pytest.mark.parametrize("style", ["report", "official", "academic"])
def test_word_builder_renders_inline_bold(style: str) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "sample.docx"
        content = "**关键词**: 柔性引才; 科技副总; 产学研协同"
        build_word_document(path, content=content, style=style)

        doc = Document(str(path))
        paragraph = doc.paragraphs[-1]
        assert "**" not in paragraph.text
        assert paragraph.text.startswith("关键词")

        bold_runs = [run.text for run in paragraph.runs if run.bold]
        plain_runs = [run.text for run in paragraph.runs if not run.bold]

        assert any("关键词" in text for text in bold_runs)
        assert any("柔性引才" in text for text in plain_runs)


def test_word_builder_renders_bold_in_bullet() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "sample.docx"
        content = "- **要点一**: 说明文字"
        build_word_document(path, content=content, style="report")

        doc = Document(str(path))
        paragraph = doc.paragraphs[-1]
        bold_runs = [run.text for run in paragraph.runs if run.bold]
        assert any("要点一" in text for text in bold_runs)
        assert "**" not in paragraph.text


def test_word_builder_leaves_unclosed_bold_markers_as_text() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "sample.docx"
        content = "**未闭合加粗"
        build_word_document(path, content=content, style="report")

        doc = Document(str(path))
        paragraph = doc.paragraphs[-1]
        assert paragraph.text == "**未闭合加粗"
        assert not any(run.bold for run in paragraph.runs)


def test_heading_font_level_4_uses_h4_preset() -> None:
    preset = WORD_PRESETS["official"]
    h3_font, h3_size = _heading_font(preset, 3)
    h4_font, h4_size = _heading_font(preset, 4)
    h6_font, h6_size = _heading_font(preset, 6)

    assert (h3_font, h3_size) == (preset.h3_font, preset.h3_size_pt)
    assert (h4_font, h4_size) == (preset.h4_font, preset.h4_size_pt)
    assert (h6_font, h6_size) == (preset.h4_font, preset.h4_size_pt)
    assert h4_size < h3_size


def test_official_h4_heading_renders_with_h4_size() -> None:
    preset = WORD_PRESETS["official"]
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "sample.docx"
        build_word_document(path, content="（1）四级标题", style="official")

        doc = Document(str(path))
        paragraph = doc.paragraphs[-1]
        assert paragraph.text == "四级标题"
        assert paragraph.runs[0].font.size.pt == preset.h4_size_pt


def test_bullet_falls_back_when_list_style_missing() -> None:
    from src.tools.builtin.file_io.word_builder import _add_bullet_paragraph

    class FakeParagraph:
        def __init__(self) -> None:
            self.runs: list[str] = []

        def add_run(self, text: str):
            self.runs.append(text)
            return self

    plain = FakeParagraph()

    class FakeDocument:
        def add_paragraph(self, text: str = "", style: str | None = None):
            if style == "List Bullet":
                raise KeyError("style")
            return plain

    paragraph = _add_bullet_paragraph(FakeDocument(), "列表项")
    assert paragraph is plain
    assert plain.runs == ["• "]


def test_template_preserves_margins_and_skips_preset_body_font() -> None:
    from docx import Document as DocxDocument
    from docx.shared import Mm, Pt

    preset = WORD_PRESETS["official"]
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        template = tmp_path / "template.docx"
        output = tmp_path / "out.docx"

        template_doc = DocxDocument()
        template_doc.sections[0].left_margin = Mm(15)
        template_doc.save(str(template))

        build_word_document(
            output,
            content="模板正文",
            style="official",
            template_path=template,
        )

        doc = Document(str(output))
        assert abs(int(doc.sections[0].left_margin) - int(Mm(15))) < 1000
        assert doc.sections[0].left_margin != Mm(preset.left_margin_mm)
        body_run = doc.paragraphs[-1].runs[0]
        assert body_run.font.size != Pt(preset.body_size_pt)
