"""Tests for /api/workspace/file and /api/workspace/file-raw."""

from __future__ import annotations

import socket
from pathlib import Path
from types import SimpleNamespace

import pytest
from aiohttp import web
from aiohttp.test_utils import TestClient, TestServer

from src.ui.web_server import WebServer
from tests.helpers import make_test_coara


def _free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def _make_server(workspace: Path) -> WebServer:
    server = WebServer(make_test_coara(workspace), workspace_dir=workspace, port=_free_port())
    server.root = SimpleNamespace(foreground_coara=SimpleNamespace(workspace_dir=workspace))
    return server


def _make_app(server: WebServer) -> web.Application:
    app = web.Application()
    app.router.add_get("/api/workspace/file", server._handle_workspace_file)
    app.router.add_get("/api/workspace/file-raw", server._handle_workspace_file_raw)
    return app


@pytest.mark.asyncio
async def test_text_file_relative_path(tmp_path: Path) -> None:
    (tmp_path / "hello.py").write_text("print('hi')\n", encoding="utf-8")
    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get("/api/workspace/file", params={"token": server.auth_token, "path": "hello.py"})
        assert resp.status == 200
        data = await resp.json()
        assert data["type"] == "text"
        assert data["name"] == "hello.py"
        assert data["language"] == "py"
        assert "print('hi')" in data["content"]
        assert data["truncated"] is False


@pytest.mark.asyncio
async def test_absolute_path_outside_workspace(tmp_path: Path, tmp_path_factory: pytest.TempPathFactory) -> None:
    outside = tmp_path_factory.mktemp("outside")
    target = outside / "notes.md"
    target.write_text("# 标题\n正文\n", encoding="utf-8")
    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get("/api/workspace/file", params={"token": server.auth_token, "path": str(target)})
        assert resp.status == 200
        data = await resp.json()
        assert data["type"] == "text"
        assert data["name"] == "notes.md"
        assert data["language"] == "md"
        assert "正文" in data["content"]


@pytest.mark.asyncio
async def test_relative_path_escape_rejected(tmp_path: Path) -> None:
    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get(
            "/api/workspace/file",
            params={"token": server.auth_token, "path": "../../../../outside.txt"},
        )
        assert resp.status == 403


@pytest.mark.asyncio
async def test_office_docx_extracted(tmp_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("月报", level=1)
    doc.add_paragraph("正文内容")
    doc.save(str(tmp_path / "report.docx"))

    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get("/api/workspace/file", params={"token": server.auth_token, "path": "report.docx"})
        assert resp.status == 200
        data = await resp.json()
        assert data["type"] == "office"
        assert data["name"] == "report.docx"
        assert "# 月报" in data["content"]
        assert "正文内容" in data["content"]
        assert data["size"] > 0
        assert data["truncated"] is False


@pytest.mark.asyncio
async def test_office_extract_failure_returns_500(tmp_path: Path) -> None:
    (tmp_path / "broken.docx").write_bytes(b"not a real docx")
    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get("/api/workspace/file", params={"token": server.auth_token, "path": "broken.docx"})
        assert resp.status == 500
        assert "error" in await resp.json()


@pytest.mark.asyncio
async def test_binary_by_suffix(tmp_path: Path) -> None:
    (tmp_path / "clip.mp4").write_bytes(b"\x00\x00\x00\x18ftypmp42")
    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get("/api/workspace/file", params={"token": server.auth_token, "path": "clip.mp4"})
        assert resp.status == 200
        data = await resp.json()
        assert data == {
            "path": "clip.mp4",
            "name": "clip.mp4",
            "type": "binary",
            "mime": "video/mp4",
            "size": 12,
        }


@pytest.mark.asyncio
async def test_binary_sniffed_from_nul_bytes(tmp_path: Path) -> None:
    payload = b"\x00\x01\x02binary\x00stuff"
    (tmp_path / "data.bin").write_bytes(payload)
    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get("/api/workspace/file", params={"token": server.auth_token, "path": "data.bin"})
        assert resp.status == 200
        data = await resp.json()
        assert data["type"] == "binary"
        assert data["mime"] == "application/octet-stream"
        assert data["size"] == len(payload)


@pytest.mark.asyncio
async def test_file_raw_inline_and_download(tmp_path: Path) -> None:
    target = tmp_path / "报告.pdf"
    target.write_bytes(b"%PDF-1.4 fake")
    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get("/api/workspace/file-raw", params={"token": server.auth_token, "path": "报告.pdf"})
        assert resp.status == 200
        assert resp.headers["Content-Type"] == "application/pdf"
        disposition = resp.headers["Content-Disposition"]
        assert disposition.startswith("inline")
        assert "filename*=UTF-8''" in disposition
        assert (await resp.read()) == b"%PDF-1.4 fake"

        resp = await client.get(
            "/api/workspace/file-raw",
            params={"token": server.auth_token, "path": "报告.pdf", "download": "1"},
        )
        assert resp.status == 200
        disposition = resp.headers["Content-Disposition"]
        assert disposition.startswith("attachment")
        assert "filename*=UTF-8''" in disposition


@pytest.mark.asyncio
async def test_file_raw_absolute_path(tmp_path: Path, tmp_path_factory: pytest.TempPathFactory) -> None:
    outside = tmp_path_factory.mktemp("outside")
    target = outside / "song.mp3"
    target.write_bytes(b"ID3\x04\x00")
    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get("/api/workspace/file-raw", params={"token": server.auth_token, "path": str(target)})
        assert resp.status == 200
        assert resp.headers["Content-Type"] == "audio/mpeg"
        assert (await resp.read()) == b"ID3\x04\x00"


@pytest.mark.asyncio
async def test_directory_listing_and_missing_rejected(tmp_path: Path) -> None:
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "a.txt").write_text("x", encoding="utf-8")
    server = _make_server(tmp_path)
    async with TestClient(TestServer(_make_app(server))) as client:
        resp = await client.get("/api/workspace/file", params={"token": server.auth_token, "path": "sub"})
        assert resp.status == 200
        data = await resp.json()
        assert data["type"] == "directory"
        assert data["name"] == "sub"
        names = {e["name"] for e in data["entries"]}
        assert "a.txt" in names

        # file-raw 不支持目录
        resp = await client.get("/api/workspace/file-raw", params={"token": server.auth_token, "path": "sub"})
        assert resp.status == 400

        resp = await client.get(
            "/api/workspace/file", params={"token": server.auth_token, "path": "missing.txt"}
        )
        assert resp.status == 400
        resp = await client.get("/api/workspace/file", params={"token": server.auth_token})
        assert resp.status == 400
        resp = await client.get("/api/workspace/file-raw", params={"token": server.auth_token})
        assert resp.status == 400
